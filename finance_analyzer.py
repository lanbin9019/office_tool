import os
from typing import Dict, Any, List, Optional
from datetime import datetime
import re

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, numbers
from openpyxl.utils import get_column_letter

from office_tool.utils import generate_output_filename


class FinanceAnalyzer:
    """财务统计分析器类"""
    
    def __init__(self):
        self._header_font = Font(bold=True, size=11, color="FFFFFF")
        self._header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
        self._subheader_fill = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
        self._total_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        self._alignment = Alignment(horizontal='center', vertical='center')
        self._alignment_left = Alignment(horizontal='left', vertical='center')
        self._thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        self._money_format = '#,##0.00'
    
    def analyze(self, file_path: str, settings: Dict[str, str], log_callback=None) -> str:
        """
        进行财务统计分析
        
        Args:
            file_path: Excel文件路径
            settings: 分析设置
            log_callback: 日志回调函数
            
        Returns:
            输出文件路径
        """
        if log_callback:
            log_callback("正在读取财务数据...")
        
        df = self._read_finance_data(file_path)
        
        if df.empty:
            raise ValueError("无法读取财务数据，请检查文件格式")
        
        if log_callback:
            log_callback(f"读取到 {len(df)} 条数据记录")
        
        date_col = settings.get('date_col', '日期')
        income_col = settings.get('income_col', '收入')
        expense_col = settings.get('expense_col', '支出')
        group_by = settings.get('group_by', '按月份')
        
        actual_date_col = self._find_column(df, date_col, ['日期', '时间', 'date', 'Date', '时间戳'])
        actual_income_col = self._find_column(df, income_col, ['收入', '进账', '收款', 'income', 'Income', '收入金额'])
        actual_expense_col = self._find_column(df, expense_col, ['支出', '出账', '付款', 'expense', 'Expense', '支出金额'])
        
        if log_callback:
            log_callback(f"使用列: 日期={actual_date_col}, 收入={actual_income_col}, 支出={actual_expense_col}")
        
        df = self._process_finance_data(df, actual_date_col, actual_income_col, actual_expense_col)
        
        if log_callback:
            log_callback("正在进行统计分析...")
        
        summary = self._calculate_summary(df, actual_income_col, actual_expense_col)
        
        grouped_stats = self._group_by_period(df, actual_date_col, actual_income_col, actual_expense_col, group_by)
        
        trend_analysis = self._analyze_trend(df, actual_date_col, actual_income_col, actual_expense_col)
        
        output_path = generate_output_filename(file_path, "finance_report")
        
        if log_callback:
            log_callback("正在生成财务报告...")
        
        self._save_finance_report(
            output_path, summary, grouped_stats, trend_analysis, 
            group_by, actual_income_col, actual_expense_col
        )
        
        if log_callback:
            log_callback(f"财务统计完成: 总收入={summary['总收入']}, 总支出={summary['总支出']}, 净利润={summary['净利润']}")
        
        return output_path
    
    def _read_finance_data(self, file_path: str) -> pd.DataFrame:
        """读取财务数据"""
        try:
            xl = pd.ExcelFile(file_path)
            for sheet_name in xl.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
                if len(df.columns) >= 3:
                    return df
            return pd.read_excel(file_path)
        except Exception as e:
            try:
                wb = load_workbook(file_path, data_only=True)
                ws = wb.active
                data = []
                for row in ws.iter_rows(values_only=True):
                    data.append(row)
                wb.close()
                if data:
                    headers = data[0]
                    rows = data[1:]
                    return pd.DataFrame(rows, columns=headers)
                return pd.DataFrame()
            except:
                raise ValueError(f"无法读取Excel文件: {str(e)}")
    
    def _find_column(self, df: pd.DataFrame, preferred: str, alternatives: List[str]) -> str:
        """查找匹配的列名"""
        columns_lower = {col.lower(): col for col in df.columns}
        
        if preferred.lower() in columns_lower:
            return columns_lower[preferred.lower()]
        
        for alt in alternatives:
            if alt.lower() in columns_lower:
                return columns_lower[alt.lower()]
        
        for col in df.columns:
            for alt in alternatives:
                if alt.lower() in col.lower():
                    return col
        
        return str(df.columns[0]) if len(df.columns) > 0 else ''
    
    def _process_finance_data(self, df: pd.DataFrame, date_col: str, 
                               income_col: str, expense_col: str) -> pd.DataFrame:
        """处理财务数据"""
        result_df = df.copy()
        
        if date_col in result_df.columns:
            result_df[date_col] = pd.to_datetime(result_df[date_col], errors='coerce')
        
        for col in [income_col, expense_col]:
            if col in result_df.columns:
                result_df[col] = pd.to_numeric(result_df[col], errors='coerce').fillna(0)
        
        result_df['净利润'] = result_df.get(income_col, 0) - result_df.get(expense_col, 0)
        
        return result_df
    
    def _calculate_summary(self, df: pd.DataFrame, income_col: str, expense_col: str) -> Dict[str, Any]:
        """计算总体统计"""
        total_income = float(df[income_col].sum()) if income_col in df.columns else 0
        total_expense = float(df[expense_col].sum()) if expense_col in df.columns else 0
        net_profit = total_income - total_expense
        
        valid_transactions = len(df[(df[income_col] > 0) | (df[expense_col] > 0)]) if income_col in df.columns else 0
        
        avg_income = float(df[income_col][df[income_col] > 0].mean()) if income_col in df.columns and len(df[df[income_col] > 0]) > 0 else 0
        avg_expense = float(df[expense_col][df[expense_col] > 0].mean()) if expense_col in df.columns and len(df[df[expense_col] > 0]) > 0 else 0
        
        max_income = float(df[income_col].max()) if income_col in df.columns else 0
        max_expense = float(df[expense_col].max()) if expense_col in df.columns else 0
        
        profit_rate = (net_profit / total_income * 100) if total_income > 0 else 0
        
        return {
            '总收入': round(total_income, 2),
            '总支出': round(total_expense, 2),
            '净利润': round(net_profit, 2),
            '利润率(%)': round(profit_rate, 2),
            '有效交易数': valid_transactions,
            '平均收入': round(avg_income, 2),
            '平均支出': round(avg_expense, 2),
            '最大单笔收入': round(max_income, 2),
            '最大单笔支出': round(max_expense, 2),
        }
    
    def _group_by_period(self, df: pd.DataFrame, date_col: str, 
                         income_col: str, expense_col: str, 
                         group_by: str) -> pd.DataFrame:
        """按周期分组统计"""
        if date_col not in df.columns:
            return pd.DataFrame()
        
        temp_df = df.copy()
        temp_df = temp_df.dropna(subset=[date_col])
        
        if temp_df.empty:
            return pd.DataFrame()
        
        if group_by == '按月份':
            temp_df['统计周期'] = temp_df[date_col].dt.to_period('M')
        elif group_by == '按季度':
            temp_df['统计周期'] = temp_df[date_col].dt.to_period('Q')
        elif group_by == '按年份':
            temp_df['统计周期'] = temp_df[date_col].dt.to_period('Y')
        else:
            temp_df['统计周期'] = temp_df[date_col].dt.to_period('M')
        
        agg_dict = {}
        if income_col in temp_df.columns:
            agg_dict[income_col] = 'sum'
        if expense_col in temp_df.columns:
            agg_dict[expense_col] = 'sum'
        
        if not agg_dict:
            return pd.DataFrame()
        
        grouped = temp_df.groupby('统计周期').agg(agg_dict).reset_index()
        
        grouped['统计周期'] = grouped['统计周期'].astype(str)
        
        if income_col in grouped.columns and expense_col in grouped.columns:
            grouped['净利润'] = grouped[income_col] - grouped[expense_col]
        
        return grouped
    
    def _analyze_trend(self, df: pd.DataFrame, date_col: str, 
                        income_col: str, expense_col: str) -> Dict[str, Any]:
        """分析趋势"""
        if date_col not in df.columns:
            return {}
        
        temp_df = df.copy()
        temp_df = temp_df.dropna(subset=[date_col])
        
        if temp_df.empty:
            return {}
        
        temp_df = temp_df.sort_values(by=date_col)
        
        temp_df['月份'] = temp_df[date_col].dt.to_period('M')
        monthly = temp_df.groupby('月份').agg({
            income_col: 'sum' if income_col in temp_df.columns else 'count',
            expense_col: 'sum' if expense_col in temp_df.columns else 'count'
        }).reset_index()
        
        if len(monthly) < 2:
            return {'趋势分析': '数据不足，无法进行趋势分析'}
        
        monthly.columns = ['月份', '收入', '支出']
        monthly['净利润'] = monthly['收入'] - monthly['支出']
        
        income_trend = monthly['收入'].iloc[-1] - monthly['收入'].iloc[-2]
        expense_trend = monthly['支出'].iloc[-1] - monthly['支出'].iloc[-2]
        profit_trend = monthly['净利润'].iloc[-1] - monthly['净利润'].iloc[-2]
        
        return {
            '最近月份收入变化': round(income_trend, 2),
            '最近月份支出变化': round(expense_trend, 2),
            '最近月份利润变化': round(profit_trend, 2),
            '收入趋势': '上升' if income_trend > 0 else ('下降' if income_trend < 0 else '持平'),
            '支出趋势': '上升' if expense_trend > 0 else ('下降' if expense_trend < 0 else '持平'),
            '利润趋势': '上升' if profit_trend > 0 else ('下降' if profit_trend < 0 else '持平'),
        }
    
    def _save_finance_report(self, output_path: str, summary: Dict[str, Any], 
                              grouped_stats: pd.DataFrame, trend_analysis: Dict[str, Any],
                              group_by: str, income_col: str, expense_col: str):
        """保存财务报告到Excel"""
        wb = Workbook()
        
        ws_summary = wb.active
        ws_summary.title = "财务概览"
        
        row_num = 1
        
        title = ws_summary.cell(row=row_num, column=1, value="财务统计报告")
        title.font = Font(bold=True, size=16)
        ws_summary.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=4)
        row_num += 1
        
        ws_summary.cell(row=row_num, column=1, value=f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        ws_summary.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=4)
        row_num += 2
        
        ws_summary.cell(row=row_num, column=1, value="一、核心指标")
        ws_summary.cell(row=row_num, column=1).font = Font(bold=True, size=12)
        row_num += 1
        
        core_metrics = [
            ('指标', '金额'),
            ('总收入', summary.get('总收入', 0)),
            ('总支出', summary.get('总支出', 0)),
            ('净利润', summary.get('净利润', 0)),
            ('利润率(%)', summary.get('利润率(%)', 0)),
        ]
        
        for metric, value in core_metrics:
            cell1 = ws_summary.cell(row=row_num, column=1, value=metric)
            cell2 = ws_summary.cell(row=row_num, column=2, value=value)
            
            if '利润' in metric and isinstance(value, (int, float)) and value < 0:
                cell2.font = Font(color="FF0000")
            elif metric == '指标':
                cell1.font = self._header_font
                cell1.fill = self._header_fill
                cell2.font = self._header_font
                cell2.fill = self._header_fill
            
            if isinstance(value, (int, float)):
                cell2.number_format = self._money_format
            
            row_num += 1
        
        row_num += 1
        
        ws_summary.cell(row=row_num, column=1, value="二、详细统计")
        ws_summary.cell(row=row_num, column=1).font = Font(bold=True, size=12)
        row_num += 1
        
        detail_metrics = [
            ('指标', '数值'),
            ('有效交易数', summary.get('有效交易数', 0)),
            ('平均收入', summary.get('平均收入', 0)),
            ('平均支出', summary.get('平均支出', 0)),
            ('最大单笔收入', summary.get('最大单笔收入', 0)),
            ('最大单笔支出', summary.get('最大单笔支出', 0)),
        ]
        
        for metric, value in detail_metrics:
            cell1 = ws_summary.cell(row=row_num, column=1, value=metric)
            cell2 = ws_summary.cell(row=row_num, column=2, value=value)
            
            if metric == '指标':
                cell1.font = self._header_font
                cell1.fill = self._header_fill
                cell2.font = self._header_font
                cell2.fill = self._header_fill
            
            if isinstance(value, (int, float)) and '收入' in metric or '支出' in metric:
                cell2.number_format = self._money_format
            
            row_num += 1
        
        row_num += 1
        
        if trend_analysis:
            ws_summary.cell(row=row_num, column=1, value="三、趋势分析")
            ws_summary.cell(row=row_num, column=1).font = Font(bold=True, size=12)
            row_num += 1
            
            trend_headers = ['分析项', '结果']
            for col_idx, header in enumerate(trend_headers, 1):
                cell = ws_summary.cell(row=row_num, column=col_idx, value=header)
                cell.font = self._header_font
                cell.fill = self._header_fill
            row_num += 1
            
            for key, value in trend_analysis.items():
                ws_summary.cell(row=row_num, column=1, value=key)
                ws_summary.cell(row=row_num, column=2, value=value)
                row_num += 1
        
        for col in range(1, 5):
            ws_summary.column_dimensions[get_column_letter(col)].width = 20
        
        if not grouped_stats.empty:
            ws_grouped = wb.create_sheet(title=f"分组统计({group_by})")
            
            headers = list(grouped_stats.columns)
            for col_idx, header in enumerate(headers, 1):
                cell = ws_grouped.cell(row=1, column=col_idx, value=str(header))
                cell.font = self._header_font
                cell.fill = self._header_fill
                cell.alignment = self._alignment
            
            for r_idx, row in enumerate(grouped_stats.values, 2):
                for c_idx, value in enumerate(row, 1):
                    cell = ws_grouped.cell(row=r_idx, column=c_idx, value=value)
                    if isinstance(value, (int, float)):
                        cell.number_format = self._money_format
            
            for col_idx in range(1, len(headers) + 1):
                ws_grouped.column_dimensions[get_column_letter(col_idx)].width = 18
        
        wb.save(output_path)
