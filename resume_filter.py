import os
import re
from typing import List, Dict, Any, Tuple
from datetime import datetime
from collections import Counter

import docx2txt
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

import PyPDF2
import pdfplumber

from office_tool.utils import generate_output_filename


class ResumeFilter:
    """简历筛选器类"""
    
    def __init__(self):
        self._education_levels = {
            '不限': 0,
            '大专': 1,
            '本科': 2,
            '硕士': 3,
            '博士': 4
        }
        
        self._education_keywords = {
            '大专': ['大专', '专科', '高职', 'college', 'associate'],
            '本科': ['本科', '学士', '大学本科', 'bachelor', 'undergraduate'],
            '硕士': ['硕士', '研究生', 'master', 'postgraduate'],
            '博士': ['博士', 'phd', 'doctor', '博士后']
        }
        
        self._experience_patterns = [
            r'(\d+)\s*[年~-]?\s*(\d+)?\s*[个]?月',
            r'(\d+)\s*年\s*(\d+)?\s*[个]?月',
            r'(\d+)\s*[多]?年',
            r'(\d+)\s*[个]?月',
        ]
    
    def filter_resumes(self, file_paths: List[str], settings: Dict[str, Any], 
                       log_callback=None) -> Dict[str, Any]:
        """
        筛选简历
        
        Args:
            file_paths: 简历文件路径列表
            settings: 筛选设置
            log_callback: 日志回调函数
            
        Returns:
            筛选结果
        """
        if log_callback:
            log_callback(f"开始筛选 {len(file_paths)} 份简历...")
        
        matched_resumes = []
        unmatched_resumes = []
        
        keywords = settings.get('keywords', [])
        exclude_keywords = settings.get('exclude_keywords', [])
        min_education = settings.get('min_education', '不限')
        min_experience = settings.get('min_experience', 0)
        
        for i, file_path in enumerate(file_paths):
            file_name = os.path.basename(file_path)
            
            if log_callback:
                log_callback(f"正在分析: {file_name} ({i+1}/{len(file_paths)})")
            
            try:
                text = self._extract_text(file_path)
                
                if not text.strip():
                    if log_callback:
                        log_callback(f"警告: {file_name} 无法提取文本内容")
                    unmatched_resumes.append({
                        'file_name': file_name,
                        'file_path': file_path,
                        'reason': '无法提取文本内容',
                        'info': {}
                    })
                    continue
                
                info = self._extract_resume_info(text, file_name)
                
                is_matched, reason = self._check_eligibility(
                    info, keywords, exclude_keywords, min_education, min_experience
                )
                
                if is_matched:
                    matched_resumes.append({
                        'file_name': file_name,
                        'file_path': file_path,
                        'info': info
                    })
                else:
                    unmatched_resumes.append({
                        'file_name': file_name,
                        'file_path': file_path,
                        'reason': reason,
                        'info': info
                    })
                    
            except Exception as e:
                if log_callback:
                    log_callback(f"错误: 处理 {file_name} 时出错: {str(e)}")
                unmatched_resumes.append({
                    'file_name': file_name,
                    'file_path': file_path,
                    'reason': f'处理错误: {str(e)}',
                    'info': {}
                })
        
        result = {
            'matched_count': len(matched_resumes),
            'unmatched_count': len(unmatched_resumes),
            'matched_resumes': matched_resumes,
            'unmatched_resumes': unmatched_resumes,
            'settings': settings
        }
        
        if file_paths:
            report_path = self._generate_filter_report(result, file_paths[0], log_callback)
            result['report_path'] = report_path
        
        if log_callback:
            log_callback(f"筛选完成: 符合条件 {len(matched_resumes)} 份，不符合条件 {len(unmatched_resumes)} 份")
        
        return result
    
    def _extract_text(self, file_path: str) -> str:
        """提取简历文本"""
        ext = os.path.splitext(file_path)[1].lower()
        
        if ext == '.pdf':
            return self._extract_pdf_text(file_path)
        elif ext in ['.docx', '.doc']:
            return self._extract_word_text(file_path)
        else:
            try:
                with open(file_path, 'r', encoding='utf-8') as f:
                    return f.read()
            except:
                try:
                    with open(file_path, 'r', encoding='gbk') as f:
                        return f.read()
                except:
                    return ''
    
    def _extract_pdf_text(self, file_path: str) -> str:
        """提取PDF文本"""
        try:
            all_text = []
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        all_text.append(text)
            if all_text:
                return '\n'.join(all_text)
        except:
            pass
        
        try:
            all_text = []
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text = page.extract_text()
                    if text:
                        all_text.append(text)
            if all_text:
                return '\n'.join(all_text)
        except:
            pass
        
        return ''
    
    def _extract_word_text(self, file_path: str) -> str:
        """提取Word文本"""
        try:
            text = docx2txt.process(file_path)
            if text.strip():
                return text
        except:
            pass
        
        try:
            doc = Document(file_path)
            text_parts = []
            for para in doc.paragraphs:
                text_parts.append(para.text)
            return '\n'.join(text_parts)
        except:
            pass
        
        return ''
    
    def _extract_resume_info(self, text: str, file_name: str) -> Dict[str, Any]:
        """提取简历信息"""
        info = {
            'name': self._extract_name(text, file_name),
            'education': self._extract_education(text),
            'education_level': self._get_education_level(self._extract_education(text)),
            'experience_years': self._extract_experience(text),
            'skills': self._extract_skills(text),
            'contact': self._extract_contact(text),
            'email': self._extract_email(text),
            'phone': self._extract_phone(text),
            'keywords_found': [],
            'total_text_length': len(text)
        }
        
        return info
    
    def _extract_name(self, text: str, file_name: str) -> str:
        """提取姓名"""
        base_name = os.path.splitext(file_name)[0]
        base_name = re.sub(r'[_\-\d\s]+', '', base_name)
        if 2 <= len(base_name) <= 4:
            return base_name
        
        name_patterns = [
            r'姓\s*名\s*[：:]\s*([\u4e00-\u9fa5]{2,4})',
            r'([\u4e00-\u9fa5]{2,4})\s*[的]?\s*简\s*历',
            r'^([\u4e00-\u9fa5]{2,4})',
        ]
        
        for pattern in name_patterns:
            match = re.search(pattern, text)
            if match:
                return match.group(1)
        
        return '未知'
    
    def _extract_education(self, text: str) -> str:
        """提取学历"""
        text_lower = text.lower()
        
        for edu_level, keywords in self._education_keywords.items():
            for keyword in keywords:
                if keyword.lower() in text_lower:
                    return edu_level
        
        return '未知'
    
    def _get_education_level(self, education: str) -> int:
        """获取学历等级"""
        return self._education_levels.get(education, 0)
    
    def _extract_experience(self, text: str) -> float:
        """提取工作经验（年）"""
        text_lower = text.lower()
        
        experience_keywords = ['工作经验', '工作年限', '从业经验', 'work experience', 'experience']
        has_experience_section = any(kw in text_lower for kw in experience_keywords)
        
        max_years = 0
        
        for pattern in self._experience_patterns:
            matches = re.findall(pattern, text)
            for match in matches:
                if isinstance(match, tuple):
                    years = int(match[0]) if match[0] else 0
                    months = int(match[1]) if len(match) > 1 and match[1] else 0
                    total = years + months / 12
                else:
                    total = int(match) if match else 0
                
                if total > max_years and total <= 50:
                    max_years = total
        
        if '应届' in text or '应届生' in text or '毕业生' in text:
            return 0
        
        return round(max_years, 1)
    
    def _extract_skills(self, text: str) -> List[str]:
        """提取技能"""
        skills = []
        
        skill_keywords = [
            'python', 'java', 'c\\+\\+', 'c#', 'javascript', 'js', 'typescript',
            'html', 'css', 'react', 'vue', 'angular', 'node\\.js', 'php',
            'sql', 'mysql', 'oracle', 'postgresql', 'mongodb', 'redis',
            '数据分析', '数据挖掘', '机器学习', '深度学习', '人工智能', 'ai',
            '项目管理', '产品经理', '运营', '市场', '销售', '财务', '会计',
            'excel', 'word', 'ppt', 'powerpoint', 'office',
            '英语', '日语', '韩语', '法语', '德语',
            '驾驶证', '驾照', '注册会计师', 'cpa', '律师资格证',
        ]
        
        text_lower = text.lower()
        
        for skill in skill_keywords:
            if re.search(r'\b' + skill + r'\b', text_lower) or skill in text_lower:
                if skill not in skills:
                    skills.append(skill.upper() if len(skill) <= 3 else skill.title())
        
        return skills[:20]
    
    def _extract_contact(self, text: str) -> str:
        """提取联系方式"""
        contacts = []
        
        phone = self._extract_phone(text)
        if phone:
            contacts.append(f"电话: {phone}")
        
        email = self._extract_email(text)
        if email:
            contacts.append(f"邮箱: {email}")
        
        return '; '.join(contacts) if contacts else '未找到'
    
    def _extract_phone(self, text: str) -> str:
        """提取手机号"""
        phone_patterns = [
            r'1[3-9]\d{9}',
            r'\+?86[-_\s]?1[3-9]\d{9}',
            r'\(?0\d{2,3}\)?[-_\s]?\d{7,8}',
        ]
        
        for pattern in phone_patterns:
            match = re.search(pattern, text)
            if match:
                return match.group(0)
        
        return ''
    
    def _extract_email(self, text: str) -> str:
        """提取邮箱"""
        email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
        
        match = re.search(email_pattern, text)
        if match:
            return match.group(0)
        
        return ''
    
    def _check_eligibility(self, info: Dict[str, Any], 
                           keywords: List[str], exclude_keywords: List[str],
                           min_education: str, min_experience: float) -> Tuple[bool, str]:
        """检查是否符合条件"""
        reasons = []
        
        if exclude_keywords:
            text = str(info.get('skills', [])) + str(info.get('name', ''))
            for kw in exclude_keywords:
                if kw.lower() in text.lower():
                    return False, f"包含排除关键词: {kw}"
        
        if min_education != '不限':
            required_level = self._education_levels.get(min_education, 0)
            current_level = info.get('education_level', 0)
            if current_level < required_level:
                reasons.append(f"学历不满足要求（要求: {min_education}）")
        
        if min_experience > 0:
            current_exp = info.get('experience_years', 0)
            if current_exp < min_experience:
                reasons.append(f"工作经验不满足要求（要求: {min_experience}年）")
        
        if keywords:
            text = str(info.get('skills', [])) + str(info.get('name', ''))
            found_keywords = []
            for kw in keywords:
                if kw.lower() in text.lower():
                    found_keywords.append(kw)
            
            info['keywords_found'] = found_keywords
            
            if not found_keywords:
                reasons.append(f"未找到必备关键词（需要: {', '.join(keywords)}）")
        
        if reasons:
            return False, '; '.join(reasons)
        
        return True, '符合条件'
    
    def _generate_filter_report(self, result: Dict[str, Any], 
                                 sample_path: str, log_callback=None) -> str:
        """生成筛选报告"""
        output_path = generate_output_filename(sample_path, "resume_filter_result", ".docx")
        
        doc = Document()
        
        title = doc.add_paragraph()
        title_run = title.add_run("简历筛选结果报告")
        title_run.bold = True
        title_run.font.size = Pt(16)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_paragraph("=" * 60)
        
        settings = result.get('settings', {})
        settings_para = doc.add_paragraph()
        settings_para.add_run("筛选条件:").bold = True
        doc.add_paragraph(f"  必备关键词: {', '.join(settings.get('keywords', [])) or '无'}")
        doc.add_paragraph(f"  排除关键词: {', '.join(settings.get('exclude_keywords', [])) or '无'}")
        doc.add_paragraph(f"  最低学历: {settings.get('min_education', '不限')}")
        doc.add_paragraph(f"  最低工作经验: {settings.get('min_experience', 0)} 年")
        
        doc.add_paragraph("-" * 60)
        
        summary_para = doc.add_paragraph()
        summary_para.add_run("筛选结果统计:").bold = True
        doc.add_paragraph(f"  总简历数: {result['matched_count'] + result['unmatched_count']}")
        doc.add_paragraph(f"  符合条件: {result['matched_count']} 份")
        doc.add_paragraph(f"  不符合条件: {result['unmatched_count']} 份")
        
        if result['matched_count'] > 0:
            doc.add_paragraph("=" * 60)
            matched_title = doc.add_paragraph()
            matched_title.add_run("【符合条件的简历】").bold = True
            
            for i, resume in enumerate(result['matched_resumes'], 1):
                doc.add_paragraph(f"\n{i}. {resume['file_name']}")
                info = resume.get('info', {})
                doc.add_paragraph(f"   姓名: {info.get('name', '未知')}")
                doc.add_paragraph(f"   学历: {info.get('education', '未知')}")
                doc.add_paragraph(f"   工作经验: {info.get('experience_years', 0)} 年")
                doc.add_paragraph(f"   技能: {', '.join(info.get('skills', [])) or '未提取到'}")
                doc.add_paragraph(f"   联系方式: {info.get('contact', '未找到')}")
                if info.get('keywords_found'):
                    doc.add_paragraph(f"   匹配关键词: {', '.join(info['keywords_found'])}")
        
        if result['unmatched_count'] > 0:
            doc.add_paragraph("\n" + "=" * 60)
            unmatched_title = doc.add_paragraph()
            unmatched_title.add_run("【不符合条件的简历】").bold = True
            
            for i, resume in enumerate(result['unmatched_resumes'], 1):
                doc.add_paragraph(f"\n{i}. {resume['file_name']}")
                doc.add_paragraph(f"   原因: {resume.get('reason', '未知')}")
                info = resume.get('info', {})
                if info:
                    doc.add_paragraph(f"   姓名: {info.get('name', '未知')}")
                    doc.add_paragraph(f"   学历: {info.get('education', '未知')}")
                    doc.add_paragraph(f"   工作经验: {info.get('experience_years', 0)} 年")
        
        doc.save(output_path)
        
        if log_callback:
            log_callback(f"筛选报告已生成: {output_path}")
        
        return output_path
